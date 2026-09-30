import openpyxl, re, sys, uuid

sys.stdout.reconfigure(encoding='utf-8')

# 1. Load users from D1_04-08-26.xlsx
db_users = []
wb_users = openpyxl.load_workbook(r'D:\.gemini\antigravity\scratch\D1_04-08-26.xlsx', data_only=True)
sheet_users = wb_users.active
for row in list(sheet_users.iter_rows(min_row=2, values_only=True)):
    uid = row[0]
    name = row[1]
    firstName = row[8]
    lastName = row[9]
    if name:
        db_users.append({
            'id': uid,
            'name': str(name).strip(),
            'firstName': str(firstName).strip() if firstName else '',
            'lastName': str(lastName).strip() if lastName else ''
        })

# Mapping dictionary for Thai PIs to English names in the DB
thai_to_eng = {
    'ธนิยา ภู่พัฒน์': 'Thaniya Bhupath',
    'ภัททิยา ชื่นชูศิลป์': 'Patthiya Chuenchusilp',
    'ไกลตา ศรีสิงห์': 'Klaita Srisingh',
    'รวี จงคงคาวุฒิ': 'Rawee Jongkongkawutthi',
    'อินทิพร โฆษิตานุฤทธิ์': 'Inthiporn Kositanurit',
    'ก้องภพ เรียวสงวนวงศ์': 'Kongpob Reosanguanwong',
    'ปฐพงศ์ โตวิวัฒน์': 'Patapong Towiwat',
    'ดวงนภา รุ่งพิบูลโสภิษฐ์': 'Duangnapa Roongpiboonsopit',
    'ธัชชัย ศรีเสน': 'Thuchchai Srisen',
    'อดิศวร กัวตระกูล': 'Adisuan Kuatrakul',
    'เกรียงศักดิ์ อุ่นบุญธรรม': 'Kriangsak Unbuntham',
    'ธีระชัย ธรรมาธิวัฒน์': 'Theerachai Thammathiwat',
    'สุปกาณฑ์ คงสวัสดิ์': 'Supakarn Khongsawat',
    'สุธาทิพย์ พงษ์เจริญ': 'Sutatip Pongcharoen',
    'พีระพล วอง': 'Peerapon Wong',
    'ไพสิฐ โกสุม': 'Paisit Kosum',
    'วิภาพร สุหฤทดำรง': 'Wipaporn Suharitdamrong',
    'กาญจ์รวี สังข์เปรม': 'Kanrawee Sungprem',
    'ณัฐภรณ์ ท้วมใจดี': 'Nattaporn Tuamjaidee',
    'เมธิรา คำทอง': 'Metira Kamtong',
    'วัชราภรณ์ ตาบูรี': 'Watcharaporn Taburee',
    'อรรควร มหัทธนตระกูล': 'Akaworn Mahatthanatrakul',
    'ณัฐรุจน์ ชัยพุทธานุกูล': 'Nattharut Chaibhuddanugul',
    'สันติ วีรกุล': 'Santi Weerakul',
    'ปิยะเมธ ดิลกธรสกุล': 'Piyameth Dilokthornsakul',
    'ศุภชัย ลวณะสกล': 'Supachai Lawanasakol',
    'บดินทร์ บุตรธรรม': 'Bodin Butthum',
    'ศุภณา ชื่นสกุล': 'Supana Chunsakul',
    'นิสิต พูลธนะนันท์': 'Nisit Poolthananant',
    'สุดารัตน์ อิศราวิศวกุล': 'Sudarat Isaravisavakul',
    'เจนยุทธ ศรีหิรัญ': 'Jenyuth Srihirun',
    'ประยุทธ ภูวรัตนาวิวิธ': 'Prayut Phuwarattanawiwit',
    'ศรัณยา เทศประสิทธิ์': 'Saranya Teprasit',
    'ธนกร งามสวย': 'Thanakorn Ngamsouy',
    'ธนพงษ์ ขจรไตรเดช': 'Thanapong Kajorntraidet',
    'ปิติ รัตนปรีชาเวช': 'Piti Rattanapreechavech',
    'อรณิชา นาคะรัศมี': 'Ornicha Nakarasmee',
    'ณัฏฐนิชญ์ สิทธิธนาวงศ์': 'Nattanit Sittithanawong',
    'วิทวัส จิตต์ผิวงาม': 'Wittawat Jitpewngarm',
    'พัชรดา อมาตยกุล': 'Patcharada Amatyakul',
    'เนตรญา วิโรจวานิช': 'Netraya Wirojwanich',
    'ปัณฑารีย์ ดอกคำ': 'Pandaree Dokkham',
    'นิธินา ยี่สิบแสน': 'Nithina Yisipsaen',
    'ณัฐชยา โรจน์วิมลการ': 'Natchaya Rojwimolkarn',
    'สิริกาญจน์ ธนะนู': 'Sirikarn Tananoo',
    'อารีย์ หินเพชร': 'Aree Hinphet',
    'ฟ้าสินี อรุณโรจน์ปัญญา': 'Fasinee Arunrodpanya',
    'ศุภธิดา อยู่เจริญ': 'Suphathida Yucharoen',
    'มัทนพร จึงมั่นคง': 'Mattanaporn Juengmankong',
    'รัฐภูมิ วรานุสาสน์': 'Rattapoom Waranusast'
}

prefixes = ['ศ.ดร.พญ.', 'รศ.ดร.พญ.', 'ผศ.ดร.พญ.', 'ดร.พญ.', 'พญ.', 'นพ.', 'ศ.ดร.', 'รศ.ดร.', 'ผศ.ดร.', 'ดร.', 'ศ.', 'รศ.', 'ผศ.', 'อ.', 'นาง', 'นส.']

def split_thai_name(raw_name):
    # e.g., 'ผศ.นพ.วิทวัส จิตต์ผิวงาม'
    name_str = str(raw_name).strip()
    
    # Strip (ผู้ร่วม...) or description
    name_str = re.sub(r'\(.*\)', '', name_str).strip()
    name_str = re.sub(r'วิศวฯ', '', name_str).strip()
    
    title = ''
    for p in prefixes:
        if name_str.startswith(p):
            title = p
            name_str = name_str[len(p):].strip()
            break
            
    parts = name_str.split()
    first_name = parts[0] if len(parts) > 0 else ''
    last_name = ' '.join(parts[1:]) if len(parts) > 1 else ''
    
    return title, first_name, last_name

# Date Parsing Heuristics
months_th = {
    'ม.ค.': 1, 'มค.': 1, 'มค': 1,
    'ก.พ.': 2, 'กพ.': 2, 'กพ': 2,
    'มี.ค.': 3, 'มีค.': 3, 'มีค': 3,
    'เม.ย.': 4, 'เมย.': 4, 'เมย': 4,
    'พ.ค.': 5, 'พค.': 5, 'พค': 5,
    'มิ.ย.': 6, 'มิย.': 6, 'มิย': 6,
    'ก.ค.': 7, 'กค.': 7, 'กค': 7,
    'ส.ค.': 8, 'สค.': 8, 'สค': 8,
    'ก.ย.': 9, 'กย.': 9, 'กย': 9,
    'ต.ค.': 10, 'ตค.': 10, 'ตค': 10,
    'พ.ย.': 11, 'พย.': 11, 'พย': 11,
    'ธ.ค.': 12, 'ธค.': 12, 'ธค': 12
}

def parse_thai_date(date_str):
    if not date_str:
         return None
    date_str = str(date_str).strip()
    
    # Check formats like '15 มี.ค. 65'
    match = re.search(r'(\d+)\s+([^\s\d]+)\s+(\d+)', date_str)
    if match:
        d = int(match.group(1))
        m_str = match.group(2)
        y = int(match.group(3))
        
        # Convert short year to full year (พ.ศ. -> ค.ศ.)
        if y < 100:
            y += 2500
        y_ce = y - 543
        
        m = months_th.get(m_str)
        if not m:
            # Try fuzzy match of month
            for k, v in months_th.items():
                if k in m_str:
                    m = v
                    break
        if m:
            return f"{y_ce:04d}-{m:02d}-{d:02d}T00:00:00.000Z"
            
    return None

def parse_duration(dur_str):
    if not dur_str:
        return None, None
    dur_str = str(dur_str).strip()
    
    parts = dur_str.split('-')
    if len(parts) == 2:
        start = parse_thai_date(parts[0])
        end = parse_thai_date(parts[1])
        return start, end
        
    return None, None

# Parse all sheets in Excel
wb = openpyxl.load_workbook(r'D:\.gemini\antigravity\scratch\funding67-69.xlsx', data_only=True)
all_projects = []
researcher_updates = {}

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    
    header_idx = -1
    for i, r in enumerate(rows):
        if r[0] == 'ลำดับ' or r[1] == 'รหัส' or r[2] == 'โครงการ':
            header_idx = i
            break
            
    if header_idx == -1:
        continue
        
    headers = [str(h).strip() if h is not None else f'Col_{idx}' for idx, h in enumerate(rows[header_idx])]
    
    current_category = 'Unknown'
    for r in rows[header_idx + 1:]:
        if all(x is None for x in r):
            continue
        non_none = [x for x in r if x is not None]
        if len(non_none) == 1 and isinstance(non_none[0], str) and ('แหล่งทุน' in non_none[0] or 'ทุน' in non_none[0]):
            current_category = non_none[0].strip()
            continue
        if any(isinstance(x, str) and ('รวม' in x or 'รวมทั้งสิ้น' in x) for x in r):
            continue
            
        proj_title = r[2]
        if not proj_title or str(proj_title).strip() == '':
            continue
            
        row_dict = {'sheet': sheet_name, 'category': current_category}
        for idx, val in enumerate(r):
            if idx < len(headers):
                row_dict[headers[idx]] = val
        all_projects.append(row_dict)

# Generate SQL seed file
sql_lines = []
sql_lines.append("-- ================================================")
sql_lines.append("-- Seed Research Projects from Excel 2567-2569")
sql_lines.append("-- ================================================")
sql_lines.append("")

matched_count = 0
unmatched_count = 0

for idx, p in enumerate(all_projects):
    p_id = str(uuid.uuid4())
    p_code = p.get('รหัส') or f"GEN-{idx+1:04d}"
    title = p.get('โครงการ', '').replace("'", "''").replace('\n', ' ').replace('\r', ' ')
    raw_pi = str(p.get('หัวหน้าโครงการ', '')).replace('\n', ' ').replace('\r', ' ')
    
    # Split name parts
    title_th, first_th, last_th = split_thai_name(raw_pi)
    clean_pi_name = f"{first_th} {last_th}".strip()
    
    # Try mapping to English database user
    eng_name = thai_to_eng.get(clean_pi_name)
    leader_id = None
    
    if eng_name:
        for u in db_users:
            if eng_name.lower() in u['name'].lower() or u['name'].lower() in eng_name.lower():
                leader_id = u['id']
                # Queue updates for researcher profile with Thai name details
                researcher_updates[leader_id] = {
                    'titleTh': title_th,
                    'firstNameTh': first_th,
                    'lastNameTh': last_th
                }
                matched_count += 1
                break
                
    if not leader_id:
        # If unmatched, default to a guest/unknown leader to maintain referential integrity
        leader_id = 'user-1' # Default fallback
        unmatched_count += 1
        
    # Budget parsing
    try:
        budget_initial = float(p.get('งบประมาณเต็ม', 0) or 0)
    except:
        budget_initial = 0.0
        
    # Date parsing
    start_date, end_date = None, None
    if 'ระยะเวลาตามสัญญา' in p:
        start_date, end_date = parse_duration(p.get('ระยะเวลาตามสัญญา'))
    elif 'วันเริ่มต้นสัญญา' in p and 'วันสิ้นสุดสัญญา' in p:
        start_date = parse_thai_date(p.get('วันเริ่มต้นสัญญา'))
        end_date = parse_thai_date(p.get('วันสิ้นสุดสัญญา'))
        
    if not start_date:
        # Default dates based on sheet year
        if '67' in p['sheet']:
            start_date = "2023-10-01T00:00:00.000Z"
            end_date = "2024-09-30T23:59:59.000Z"
        elif '68' in p['sheet']:
            start_date = "2024-10-01T00:00:00.000Z"
            end_date = "2025-09-30T23:59:59.000Z"
        else:
            start_date = "2025-10-01T00:00:00.000Z"
            end_date = "2026-09-30T23:59:59.000Z"
            
    # Status parsing
    raw_status = str(p.get('สถานะโครงการ', 'อยู่ระหว่างดำเนินการ'))
    status = 'ONGOING'
    if 'เสร็จสิ้น' in raw_status or 'ปิดโครงการ' in raw_status or 'Completed' in raw_status:
        status = 'COMPLETED'
    elif 'สัญญา' in raw_status or 'อนุมัติ' in raw_status or 'Approved' in raw_status:
        status = 'APPROVED'
        
    sql_lines.append(f"-- Project: {p_code} (PI: {raw_pi})")
    sql_lines.append(f"INSERT INTO irResearchProject (id, title, status, budgetInitial, budgetSpent, startDate, endDate, leaderId, department, createdAt, updatedAt)")
    sql_lines.append(f"VALUES ('{p_id}', '{title}', '{status}', {budget_initial}, 0.0, '{start_date}', '{end_date}', '{leader_id}', 'Faculty of Medicine', datetime('now'), datetime('now'));")
    sql_lines.append("")

# Append researcher profile updates to SQL
sql_lines.append("-- ================================================")
sql_lines.append("-- Update Researcher Profiles with Thai Names")
sql_lines.append("-- ================================================")
sql_lines.append("")

for uid, info in researcher_updates.items():
    sql_lines.append(f"-- User ID: {uid}")
    sql_lines.append(f"UPDATE irResearcherProfile")
    sql_lines.append(f"SET titleTh = '{info['titleTh']}', firstNameTh = '{info['firstNameTh']}', lastNameTh = '{info['lastNameTh']}', updatedAt = datetime('now')")
    sql_lines.append(f"WHERE userId = '{uid}';")
    sql_lines.append("")

# Write SQL seed to file
with open(r'D:\.gemini\antigravity\scratch\iram-backend\import_projects_67-69.sql', 'w', encoding='utf-8') as f:
    f.write('\n'.join(sql_lines))

print(f"Generated SQL import file: import_projects_67-69.sql")
print(f"Matched PIs: {matched_count}, Fallbacks: {unmatched_count}")
