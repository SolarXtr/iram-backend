import os
import sys
import json
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("1. Parsing raw D1 users from wrangler log...")
log_file = r"C:\Users\tinnakornh\AppData\Roaming\xdg.config\.wrangler\logs\wrangler-2026-09-29_06-43-30_858.log"

users_dict = {}

if os.path.exists(log_file):
    with open(log_file, "r", encoding="utf-8") as f:
        log_content = f.read()
    
    start_str = '[\n  {\n    "results": ['
    start_idx = log_content.find(start_str)
    if start_idx != -1:
        # Find the delimiter after the json array
        end_idx = log_content.find('\n---', start_idx)
        if end_idx != -1:
            json_text = log_content[start_idx:end_idx].strip()
            json_data = json.loads(json_text)
            raw_rows = json_data[0].get("results", [])
            print(f"   Successfully parsed {len(raw_rows)} users from D1 query log.")
            for r in raw_rows:
                uid = r.get("id")
                if uid:
                    users_dict[uid] = dict(r)

print("2. Parsing import_iruser.sql for roles and dates...")
import_sql = r"D:\.gemini\antigravity\scratch\iram-backend\import_iruser.sql"
if os.path.exists(import_sql):
    with open(import_sql, "r", encoding="utf-8") as f:
        sql_content = f.read()
    
    pattern = re.compile(
        r"VALUES\s*\(\s*'([^']+)'\s*,\s*'([^']*)'\s*,\s*'([^']*)'\s*,\s*'([^']*)'\s*,\s*(\d+)\s*,\s*(NULL|'[^']*')\s*,\s*(NULL|'[^']*')\s*,\s*(NULL|'[^']*')\s*,\s*(NULL|'[^']*')\s*,\s*(NULL|'[^']*')\s*\)",
        re.IGNORECASE
    )
    for match in pattern.finditer(sql_content):
        uid = match.group(1)
        name = match.group(2)
        email = match.group(3)
        role = match.group(4)
        isDeleted = int(match.group(5))
        
        def clean_val(v):
            if v == "NULL" or v is None:
                return None
            return v.strip("'")
            
        title = clean_val(match.group(6))
        firstName = clean_val(match.group(7))
        lastName = clean_val(match.group(8))
        joinDate = clean_val(match.group(9))
        resignDate = clean_val(match.group(10))
        
        if uid not in users_dict:
            users_dict[uid] = {"id": uid}
            
        u = users_dict[uid]
        u["name"] = u.get("name") or name
        u["email"] = u.get("email") or email
        u["role"] = role
        u["isDeleted"] = isDeleted
        u["title"] = u.get("title") or title
        u["firstName"] = u.get("firstName") or firstName
        u["lastName"] = u.get("lastName") or lastName
        u["joinDate"] = joinDate
        u["resignDate"] = resignDate

print("3. Parsing populate_users_consolidated.sql for Thai/Eng names, shortNameEn, Scopus, ORCID...")
pop_sql = r"D:\.gemini\antigravity\scratch\iram-backend\populate_users_consolidated.sql"
if os.path.exists(pop_sql):
    with open(pop_sql, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        m_id = re.search(r'WHERE\s+"id"\s*=\s*\'([^\']+)\'', line)
        if not m_id:
            continue
        uid = m_id.group(1)
        if uid not in users_dict:
            users_dict[uid] = {"id": uid}
        u = users_dict[uid]
        
        def extract_field(field_name):
            m = re.search(rf'"{field_name}"\s*=\s*(\'([^\']*)\'|NULL)', line)
            if m:
                if m.group(1) == "NULL":
                    return None
                return m.group(2)
            return None

        u["titleTh"] = extract_field("titleTh")
        u["firstNameTh"] = extract_field("firstNameTh")
        u["lastNameTh"] = extract_field("lastNameTh")
        u["titleEn"] = extract_field("titleEn")
        u["firstNameEn"] = extract_field("firstNameEn")
        u["lastNameEn"] = extract_field("lastNameEn")
        u["shortNameEn"] = extract_field("shortNameEn")
        u["department"] = extract_field("department") or u.get("department") or "คณะแพทยศาสตร์ มหาวิทยาลัยนเรศวร"
        u["scopusAuthorId"] = extract_field("scopusAuthorId") or u.get("scopusAuthorId")
        u["orcid"] = extract_field("orcid") or u.get("orcid")
        u["wosResearcherId"] = extract_field("wosResearcherId") or u.get("wosResearcherId")
        u["aliasesJson"] = extract_field("aliasesJson")

print(f"Total merged users: {len(users_dict)}")

# Known specific administrative roles
admin_emails = {
    "tinnakornh@nu.ac.th": {
        "role": "admin",
        "rolesJson": '["researcher", "coordinator", "admin"]',
        "academicPosition": "เจ้าหน้าที่วิจัย",
        "administrativePosition": "หัวหน้าหน่วยบริหารและจัดการงานวิจัย",
        "phone": "5588",
        "bankName": "ธนาคารกรุงศรีอยุธยา สาขามหาวิทยาลัยนเรศวร",
        "bankAccountNo": "346-1-00888-9",
        "idCardNo": "1-6500-00888-99-0"
    },
    "somkiat.r@iram.edu": {
        "role": "researcher",
        "rolesJson": '["researcher"]',
        "academicPosition": "ศาสตราจารย์",
        "titleTh": "ศ.ดร.",
        "firstNameTh": "สมเกียรติ",
        "lastNameTh": "รักเรียน"
    },
    "wipa.j@iram.edu": {
        "role": "researcher",
        "rolesJson": '["researcher"]',
        "academicPosition": "อาจารย์ / นักวิจัย",
        "titleTh": "ดร.",
        "firstNameTh": "วิภา",
        "lastNameTh": "จิตวิทยา"
    },
    "wandee.w@iram.edu": {
        "role": "coordinator",
        "rolesJson": '["coordinator"]',
        "academicPosition": "เจ้าหน้าที่งานวิจัย",
        "titleTh": "คุณ",
        "firstNameTh": "วันดี",
        "lastNameTh": "ทำงานดี"
    },
    "songpol.s@iram.edu": {
        "role": "executive",
        "rolesJson": '["executive"]',
        "academicPosition": "รองศาสตราจารย์",
        "administrativePosition": "รองคณบดีฝ่ายวิจัย",
        "titleTh": "รศ.นพ.",
        "firstNameTh": "ทรงพล",
        "lastNameTh": "บริหาร"
    }
}

for uid, u in users_dict.items():
    email = (u.get("email") or "").lower().strip()
    if email in admin_emails:
        u.update(admin_emails[email])
    
    # Defaults
    if not u.get("role"):
        u["role"] = "RESEARCHER"
    if not u.get("rolesJson"):
        clean_role = (u.get("role") or "researcher").lower()
        u["rolesJson"] = json.dumps([clean_role])
    if not u.get("academicPosition"):
        pos = u.get("titleTh") or u.get("title") or ""
        if "ศ." in pos or "ศาสตราจารย์" in pos:
            u["academicPosition"] = "ศาสตราจารย์"
        elif "รศ." in pos or "รองศาสตราจารย์" in pos:
            u["academicPosition"] = "รองศาสตราจารย์"
        elif "ผศ." in pos or "ผู้ช่วยศาสตราจารย์" in pos:
            u["academicPosition"] = "ผู้ช่วยศาสตราจารย์"
        elif "อ." in pos or "อาจารย์" in pos:
            u["academicPosition"] = "อาจารย์"
        else:
            u["academicPosition"] = "อาจารย์ / นักวิจัย"
    if not u.get("bankName"):
        u["bankName"] = "ธนาคารกรุงศรีอยุธยา สาขามหาวิทยาลัยนเรศวร"

# Convert to DataFrame with ordered columns
cols = [
    "id",
    "name",
    "email",
    "titleTh",
    "firstNameTh",
    "lastNameTh",
    "titleEn",
    "firstNameEn",
    "lastNameEn",
    "shortNameEn",
    "academicPosition",
    "administrativePosition",
    "department",
    "role",
    "rolesJson",
    "scopusAuthorId",
    "orcid",
    "wosResearcherId",
    "aliasesJson",
    "phone",
    "bankName",
    "bankAccountNo",
    "idCardNo",
    "isDeleted",
    "joinDate",
    "resignDate"
]

header_labels = [
    "ID (รหัสผู้ใช้)",
    "Full Name (ชื่อ-นามสกุล)",
    "Email (อีเมลมหาวิทยาลัย)",
    "Title (ไทย)",
    "First Name (ชื่อไทย)",
    "Last Name (นามสกุลไทย)",
    "Title (Eng)",
    "First Name (Eng)",
    "Last Name (Eng)",
    "Short Name (ชื่อย่อสากล เช่น Srisingh K.)",
    "Academic Position (ตำแหน่งวิชาการ)",
    "Administrative Position (ตำแหน่งบริหาร)",
    "Department (ภาควิชา/สังกัด)",
    "Role (สิทธิ์หลัก)",
    "Roles JSON (สิทธิ์ทั้งหมด)",
    "Scopus Author ID",
    "ORCID iD",
    "WoS ResearcherID",
    "Aliases JSON (นามแฝง/ชื่อเดิม)",
    "Phone (เบอร์ติดต่อ)",
    "Bank Name (ธนาคาร)",
    "Bank Account No (เลขที่บัญชี)",
    "ID Card No (เลขบัตรประชาชน)",
    "Is Deleted (0=Active, 1=Deleted)",
    "Join Date (วันเริ่มงาน)",
    "Resign Date (วันลาออก)"
]

rows_list = []
def sort_key(item):
    u = item[1]
    r = (u.get("role") or "").lower()
    priority = 5
    if r == "admin": priority = 1
    elif r in ["coordinator", "staff"]: priority = 2
    elif r == "finance": priority = 3
    elif r == "executive": priority = 4
    return (priority, u.get("firstNameTh") or u.get("name") or "")

sorted_users = sorted(users_dict.items(), key=sort_key)

for uid, u in sorted_users:
    row_data = [u.get(col) for col in cols]
    rows_list.append(row_data)

# Output paths
out_xlsx_scratch = r"D:\.gemini\antigravity\scratch\irUser_export_2026-09-29.xlsx"
out_xlsx_artifact = r"C:\Users\tinnakornh\.gemini\antigravity\brain\b31db261-266f-44a9-bf40-fb1f964ad43b\irUser_export.xlsx"

print(f"4. Generating Excel workbook with openpyxl...")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "irUser"

# Styling definitions
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy
header_font = Font(name="TH Sarabun New", size=14, bold=True, color="FFFFFF")
data_font = Font(name="TH Sarabun New", size=13)
data_font_bold = Font(name="TH Sarabun New", size=13, bold=True)
thin_border = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB")
)
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

# Write Technical Column Names (Row 1)
tech_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
tech_font = Font(name="Consolas", size=10, color="94A3B8", bold=True)
for col_idx, col_name in enumerate(cols, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.fill = tech_fill
    cell.font = tech_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Write Human-Readable Headers (Row 2)
for col_idx, label in enumerate(header_labels, start=1):
    cell = ws.cell(row=2, column=col_idx, value=label)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

ws.row_dimensions[1].height = 20
ws.row_dimensions[2].height = 36

# Write Data (Rows 3+)
center_cols = {1, 4, 7, 10, 14, 16, 17, 18, 20, 24, 25, 26}

for row_idx, row_values in enumerate(rows_list, start=3):
    is_even = (row_idx % 2 == 0)
    ws.row_dimensions[row_idx].height = 24
    for col_idx, val in enumerate(row_values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
        cell.font = data_font
        cell.border = thin_border
        
        if is_even:
            cell.fill = zebra_fill
            
        if col_idx in center_cols:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Auto-adjust column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        v = str(cell.value or "")
        char_len = sum(1.5 if ord(c) > 128 else 1.0 for c in v)
        if char_len > max_len:
            max_len = char_len
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Freeze top 2 header rows
ws.freeze_panes = "A3"

# Save workbooks
wb.save(out_xlsx_scratch)
wb.save(out_xlsx_artifact)
print(f"Successfully exported {len(rows_list)} users to:")
print(f"  - {out_xlsx_scratch}")
print(f"  - {out_xlsx_artifact}")
